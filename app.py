from flask import Flask, render_template, g, request, send_file
import sqlite3
import contextlib
import configparser
import os
import pandas as pd
from datetime import datetime
import io
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# 讀取 config.ini 中的 SQLite 檔案路徑
config = configparser.ConfigParser()
config.read("config.ini", encoding="utf-8")

DB_FILE = config.get("DATABASE", "db_file", fallback="resume.db")
if not os.path.isabs(DB_FILE):
    DB_FILE = os.path.join(os.path.dirname(__file__), DB_FILE)


def get_db():
    db = getattr(g, "_database", None)
    if db is None:
        db = g._database = sqlite3.connect(DB_FILE, timeout=30.0, check_same_thread=False)
        db.row_factory = sqlite3.Row
    return db


@app.teardown_appcontext
def close_connection(exception):
    db = getattr(g, "_database", None)
    if db is not None:
        db.close()


@contextlib.contextmanager
def get_cursor(db):
    cur = db.cursor()
    try:
        yield cur
    finally:
        cur.close()


import re


def parse_it_months(s):
    if not s or s.lower() == "none" or s == "":
        return -1
    total = 0
    y = re.search(r"(\d+)年", s)
    m = re.search(r"(\d+)個月", s)
    if y:
        total += int(y.group(1)) * 12
    if m:
        total += int(m.group(1))
    return total


@app.route("/")
def index():
    # 獲取篩選參數 (多選)
    selected_statuses = request.args.getlist("status")
    selected_invites = request.args.getlist("invite")
    selected_ages = request.args.getlist("age")
    selected_it_exps = request.args.getlist("it_exp")
    selected_edus = request.args.getlist("edu")
    selected_schools = request.args.getlist("school")
    selected_majors = request.args.getlist("major")

    highlight_query = request.args.get("highlight", "").strip()
    min_score = request.args.get("min_score", 0, type=int)
    max_score = request.args.get("max_score", 100, type=int)
    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 250, type=int)

    db = get_db()
    with get_cursor(db) as cur:
        # --- 構建 SQL 篩選條件 (簡單欄位) ---
        where_clauses = []
        params = []

        if selected_statuses:
            where_clauses.append(
                f"is_screened IN ({','.join(['?']*len(selected_statuses))})"
            )
            params.extend(selected_statuses)

        if selected_invites:
            invite_conditions = []
            if "是" in selected_invites:
                invite_conditions.append("interview_state = '是'")
            if "否" in selected_invites:
                invite_conditions.append("(interview_state = '否' OR interview_state = '' OR interview_state IS NULL)")
            
            if invite_conditions:
                where_clauses.append(f"({' OR '.join(invite_conditions)})")

        if selected_edus:
            # 學歷可能是包含關係
            edu_clause = " OR ".join(
                ["highest_edu_level LIKE ?" for _ in selected_edus]
            )
            where_clauses.append(f"({edu_clause})")
            params.extend([f"%{e}%" for e in selected_edus])

        if selected_schools:
            where_clauses.append(
                f"highest_edu_school IN ({','.join(['?']*len(selected_schools))})"
            )
            params.extend(selected_schools)

        if selected_majors:
            where_clauses.append(
                f"highest_edu_major IN ({','.join(['?']*len(selected_majors))})"
            )
            params.extend(selected_majors)

        if highlight_query:
            where_clauses.append(
                "(autobiography LIKE ? OR work_experience LIKE ? OR key_highlights LIKE ?)"
            )
            params.extend([f"%{highlight_query}%"] * 3)

        if min_score > 0 or max_score < 100:
            where_clauses.append("score BETWEEN ? AND ?")
            params.extend([min_score, max_score])

        where_sql = ""
        if where_clauses:
            where_sql = "WHERE " + " AND ".join(where_clauses)

        # 1. 執行初步篩選
        cur.execute(
            f"SELECT *, rowid FROM curriculum_vitae {where_sql} ORDER BY rowid DESC",
            params,
        )
        all_matches = [dict(r) for r in cur.fetchall()]

        # 2. Python 處理複雜篩選 (Age, IT Experience)
        filtered_resumes = []
        for r in all_matches:
            # 年齡篩選
            if selected_ages:
                try:
                    age_val = int(r.get("age", 0))
                except:
                    age_val = 0
                match_age = False
                for range_str in selected_ages:
                    age_min, age_max = map(int, range_str.split("-"))
                    if age_min <= age_val <= age_max:
                        match_age = True
                        break
                if not match_age:
                    continue

            # IT 年資篩選
            if selected_it_exps:
                it_months = parse_it_months(r.get("it_work_years", ""))
                match_exp = False
                for range_str in selected_it_exps:
                    if range_str == "None":
                        if it_months == -1:
                            match_exp = True
                    else:
                        min_y, max_y = map(float, range_str.split("-"))
                        years = it_months / 12.0
                        if years >= min_y and years < max_y:
                            match_exp = True
                    if match_exp:
                        break
                if not match_exp:
                    continue

            filtered_resumes.append(r)

        # 2.1 執行排序 (Sort all filtered results)
        sort_by = request.args.get("sort_by")
        order = request.args.get("order", "asc")

        if sort_by:
            reverse = order == "desc"

            def sort_key(r):
                val = r.get(sort_by)
                if val is None:
                    return ""

                # 數值類
                if sort_by in ["score", "age", "code"]:
                    try:
                        return int(val)
                    except:
                        return 0

                # 年資類
                if sort_by in ["total_work_years", "it_work_years"]:
                    return parse_it_months(str(val))

                # 學歷類
                if sort_by == "highest_edu_level":
                    edu_weights = {
                        "博士": 5,
                        "碩士": 4,
                        "大學": 3,
                        "專科": 2,
                        "高中": 1,
                    }
                    return next((w for k, w in edu_weights.items() if k in str(val)), 0)

                return str(val)

            filtered_resumes.sort(key=sort_key, reverse=reverse)

        # 3. 分頁邏輯 (一頁 250 筆)
        total_resumes = len(filtered_resumes)
        total_pages = (total_resumes + per_page - 1) // per_page
        if page < 1:
            page = 1
        if page > total_pages and total_pages > 0:
            page = total_pages

        offset = (page - 1) * per_page
        resumes = filtered_resumes[offset : offset + per_page]

        # --- 獲取所有篩選選單的動態選項 (不隨篩選變動，保持完整清單) ---
        cur.execute(
            'SELECT DISTINCT highest_edu_level FROM curriculum_vitae WHERE highest_edu_level IS NOT NULL AND highest_edu_level != ""'
        )
        edu_rows = cur.fetchall()
        edu_weights = {"博士": 5, "碩士": 4, "大學": 3, "專科": 2, "高中": 1}
        sorted_levels = sorted(
            [r["highest_edu_level"] for r in edu_rows],
            key=lambda x: next((w for k, w in edu_weights.items() if k in x), 0),
        )

        cur.execute(
            'SELECT DISTINCT highest_edu_school FROM curriculum_vitae WHERE highest_edu_school IS NOT NULL AND highest_edu_school != ""'
        )
        sorted_schools = sorted(
            [r["highest_edu_school"] for r in cur.fetchall()],
            key=lambda s: (0 if any("\u4e00" <= c <= "\u9fff" for c in s) else 1, s),
        )

        cur.execute(
            'SELECT DISTINCT highest_edu_major FROM curriculum_vitae WHERE highest_edu_major IS NOT NULL AND highest_edu_major != ""'
        )
        sorted_majors = sorted(
            [r["highest_edu_major"] for r in cur.fetchall()],
            key=lambda s: (0 if any("\u4e00" <= c <= "\u9fff" for c in s) else 1, s),
        )

        cur.execute(
            'SELECT DISTINCT is_screened FROM curriculum_vitae WHERE is_screened IS NOT NULL AND is_screened != ""'
        )
        raw_statuses = [r["is_screened"] for r in cur.fetchall()]
        if "未儲存" not in raw_statuses:
            raw_statuses.append("未儲存")
        sorted_statuses = sorted(
            raw_statuses, key=lambda s: (0 if s == "未儲存" else 1, s)
        )

        # --- 新增：獲取全量 學校-科系 對照表 ---
        cur.execute(
            'SELECT DISTINCT highest_edu_school, highest_edu_major FROM curriculum_vitae WHERE highest_edu_school != "" AND highest_edu_major != ""'
        )
        sm_rows = cur.fetchall()
        school_major_map = {}
        for row in sm_rows:
            s, m = row["highest_edu_school"], row["highest_edu_major"]
            if s not in school_major_map:
                school_major_map[s] = []
            school_major_map[s].append(m)

        # 次高學歷也納入考慮
        cur.execute(
            'SELECT DISTINCT second_edu_school, second_edu_major FROM curriculum_vitae WHERE second_edu_school != "" AND second_edu_major != ""'
        )
        sm_rows2 = cur.fetchall()
        for row in sm_rows2:
            s, m = row["second_edu_school"], row["second_edu_major"]
            if s not in school_major_map:
                school_major_map[s] = []
            if m not in school_major_map[s]:
                school_major_map[s].append(m)

    return render_template(
        "index.html",
        resumes=resumes,
        edu_levels=sorted_levels,
        schools=sorted_schools,
        majors=sorted_majors,
        screen_statuses=sorted_statuses,
        channels=[],
        school_major_map=school_major_map,
        total_resumes=total_resumes,
        page=page,
        total_pages=total_pages,
        per_page=per_page,
        # 回傳當前篩選狀態
        selected_filters={
            "status": selected_statuses,
            "invite": selected_invites,
            "age": selected_ages,
            "it_exp": selected_it_exps,
            "edu": selected_edus,
            "school": selected_schools,
            "major": selected_majors,
            "channel": [],
            "highlight": highlight_query,
            "min_score": min_score,
            "max_score": max_score,
            "date": "",
            "per_page": per_page,
        },
    )


@app.route("/export", methods=["POST"])
def export_excel():
    selected_ids = request.form.getlist("selected_resumes")
    if not selected_ids:
        return "未選擇任何履歷", 400

    db = get_db()
    with get_cursor(db) as cur:
        placeholders = ",".join(["?"] * len(selected_ids))
        query = f"SELECT * FROM curriculum_vitae WHERE code IN ({placeholders})"
        cur.execute(query, selected_ids)
        rows = [dict(r) for r in cur.fetchall()]

    if not rows:
        return "找不到對應的履歷資料", 404

    df = pd.DataFrame(rows)

    # 依照使用者要求：不要將 Flask 中的狀態（日期）輸出，保持與 main.py 一致
    # 將資料重設為原始預設值
    if "is_screened" in df.columns:
        df["is_screened"] = "未儲存"

    # 欄位重新命名與排序
    column_mapping = {
        "interview_state": "邀約面試狀態",
        "is_screened": "是否篩選",
        "name": "姓名",
        "code": "代碼",
        "age": "年齡",
        "total_work_years": "總工作年資",
        "it_work_years": "資訊類工作年資",
        "autobiography": "自傳",
        "work_experience": "工作經歷",
        "highest_edu_level": "最高學歷",
        "highest_edu_school": "最高學歷學校",
        "highest_edu_major": "最高學歷科系",
        "highest_edu_period": "最高學歷期間",
        "second_edu_school": "次高學歷學校",
        "second_edu_major": "次高學歷科系",
        "tech_skills": "資訊類相關專長",
        "key_highlights": "履歷亮點",
        "notes": "備註",
        "score_data_eng": "資料工程能力評估",
        "score_db_sql": "資料庫與 SQL 能力評估",
        "score_ai_app": "AI 應用能力評估",
        "score_rag_sys": "RAG 與系統整合能力評估",
        "score_deployment": "平台部署與工程能力評估",
        "score": "總分",
        "reason": "理由"
    }

    # 定義欄位順序並重新命名
    columns_order = [c for c in column_mapping.keys() if c in df.columns]
    df = df[columns_order].rename(columns=column_mapping)

    # 建立 Excel 檔案
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="履歷分析")
        workbook = writer.book
        worksheet = writer.sheets["履歷分析"]
        header_font = Font(bold=True)

        # 定義需要固定寬度並換行的欄位
        fixed_width_columns = [
            "自傳",
            "工作經歷",
            "資訊類相關專長",
            "履歷亮點",
            "備註",
            "理由"
        ]
        fixed_width_val = 80

        # 定義評分細節欄位
        score_detail_columns = [
            "資料工程能力評估",
            "資料庫與 SQL 能力評估",
            "AI 應用能力評估",
            "RAG 與系統整合能力評估",
            "平台部署與工程能力評估",
        ]
        score_detail_val = 40

        # 設定欄寬與格式 (與 main.py 一致)
        for i, col_name in enumerate(df.columns, 1):
            col_letter = get_column_letter(i)

            if col_name in fixed_width_columns:
                worksheet.column_dimensions[col_letter].width = fixed_width_val
                for cell in worksheet[col_letter]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

            elif col_name in score_detail_columns:
                worksheet.column_dimensions[col_letter].width = score_detail_val
                for cell in worksheet[col_letter]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

            elif col_name == "總分":
                # 評分欄位：縮短寬度並置中
                worksheet.column_dimensions[col_letter].width = 10
                for cell in worksheet[col_letter]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            else:
                # 自動調整寬度
                max_length = 0
                cells = [worksheet[f"{col_letter}1"]] + list(worksheet[col_letter])[1:]
                for cell in cells:
                    try:
                        val = cell.value
                        if val:
                            str_val = str(val)
                            local_len = 0
                            for char in str_val:
                                if ord(char) > 127:
                                    local_len += 2
                                else:
                                    local_len += 1
                            if local_len > max_length:
                                max_length = local_len
                    except:
                        pass

                adjusted_width = max_length + 2
                worksheet.column_dimensions[col_letter].width = adjusted_width
                for cell in worksheet[col_letter]:
                    if not cell.alignment.wrap_text:
                        cell.alignment = Alignment(vertical="center")

            # 標題列統一格式
            header_cell = worksheet[f"{col_letter}1"]
            header_cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            header_cell.font = header_font

    output.seek(0)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"104_Resume_{timestamp}.xlsx"

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/mark_viewed", methods=["POST"])
def mark_viewed():
    data = request.get_json()
    selected_ids = data.get("ids", [])
    new_status = data.get("status", "已儲存")  # 允許動態傳入狀態

    # 如果狀態是「已儲存」，改為存入當天的 YYYYMMDD
    if new_status == "已儲存":
        new_status = datetime.now().strftime("%Y%m%d")

    if not selected_ids:
        return {"status": "error", "message": "No IDs provided"}, 400

    db = get_db()
    try:
        with get_cursor(db) as cur:
            placeholders = ", ".join(["?"] * len(selected_ids))
            sql = f"UPDATE curriculum_vitae SET is_screened = ? WHERE code IN ({placeholders})"
            # 將 status 作為第一個參數傳入
            cur.execute(sql, [new_status] + selected_ids)
        db.commit()
        return {"status": "success"}

    except Exception as e:
        return {"status": "error", "message": str(e)}, 500


@app.route("/mark_invited", methods=["POST"])
def mark_invited():
    data = request.get_json()
    code = data.get("id")
    status = data.get("status", "是")

    if not code:
        return {"status": "error", "message": "No ID provided"}, 400

    db = get_db()
    try:
        with get_cursor(db) as cur:
            cur.execute(
                "UPDATE curriculum_vitae SET interview_state = ? WHERE code = ?",
                (status, code),
            )
        db.commit()
        return {"status": "success"}
    except Exception as e:
        return {"status": "error", "message": str(e)}, 500


@app.route("/mark_read", methods=["POST"])
def mark_read():
    data = request.get_json()
    code = data.get("id")
    status = data.get("status", 1)  # 預設標記為已讀(1)

    if not code:
        return {"status": "error", "message": "No ID provided"}, 400

    db = get_db()
    try:
        with get_cursor(db) as cur:
            cur.execute(
                "UPDATE curriculum_vitae SET review_status = ? WHERE code = ?",
                (status, code),
            )
        db.commit()
        return {"status": "success"}
    except Exception as e:
        return {"status": "error", "message": str(e)}, 500


if __name__ == "__main__":
    app.run(debug=True, port=5000)
